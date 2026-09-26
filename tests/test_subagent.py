# -*- coding: utf-8 -*-
"""
Субагентская сделка (через салон РФ, с 26.09.2026): полный комплект из шести
документов для салона-ООО и салона-ИП, конечный покупатель — физлицо и ИП.
Проверяется, что не осталось {{…}}, в ДКП покупатель — клиент, а не салон,
в счёте суммы числами, и что прямая сделка этого не замечает.
"""
import asyncio
import re

import openpyxl
import pytest

import doc_builder
import salon
from doc_builder import DocumentBuilder, order_amount, SubagentSetupError
from tests.test_documents import DEAL, NUMBER, PRICE, PCT, text_of, _as_num
import money

OOO = {
    "org_form": "ООО", "name_short": "ООО «Автомир»",
    "name_full": "Общество с ограниченной ответственностью «Автомир»",
    "inn": "7701234567", "kpp": "770101001", "ogrn": "1157746000000",
    "address": "г. Москва, ул. Автомобильная, д. 1",
    "signer_position": "Генеральный директор",
    "signer_position_gen": "Генерального директора",
    "signer_name": "Сидоров Сидор Сидорович", "signer_name_gen": "Сидорова Сидора Сидоровича",
    "signer_basis": "Устава",
    "account": "40702810000000000001", "bank_name": "ПАО Сбербанк",
    "bank_bic": "044525225", "bank_corr": "30101810400000000225",
}
IP = {
    "inn": "770112345678", "ogrn": "315774600000000",
    "signer_name": "Петров Пётр Петрович",
    "address": "г. Москва, ул. Предпринимателей, д. 5",
    "account": "40802810000000000002", "bank_name": "АО «Т-Банк»",
    "bank_bic": "044525974", "bank_corr": "30101810145250000974",
}


def sub_deal(card, **extra):
    c = salon.save(card)
    d = {**DEAL, "deal_type": "Субагент", "salon": salon.journal_value(c),
         "salon_contract": "№ 45 от 12.09.2026"}
    d["cash_amount"] = f"{order_amount(PRICE, 82.00, NUMBER):.2f}".replace(".", ",")
    d["Сумма выдана (USD)"] = f"{order_amount(PRICE, 82.80, NUMBER):.2f}".replace(".", ",")
    d.update(extra)
    return d


def build_all(data):
    b = DocumentBuilder()

    async def run():
        return {
            "contract": await b.build_contract(data, NUMBER, "01.07.2026", PCT),
            "dkp":      await b.build_dkp(data, NUMBER, "01.07.2026"),
            "invoice":  await b.build_invoice(data, NUMBER, "01.07.2026", PCT),
            "receipt":  await b.build_receipt(data, NUMBER, "01.07.2026", "05.07.2026", PCT,
                                              settlement_price=PRICE),
            "act":      await b.build_act(data, NUMBER, "01.07.2026", "05.07.2026", PCT,
                                          settlement_price=PRICE),
            "report":   await b.build_report(data, NUMBER, "01.07.2026", "05.07.2026",
                                             "01.07.2026", "05.07.2026", PCT, settlement_price=PRICE),
        }
    return asyncio.run(run())


def texts(paths):
    out = {}
    for k, p in paths.items():
        if p.endswith(".xlsx"):
            ws = openpyxl.load_workbook(p).active
            out[k] = "\n".join(str(c.value) for r in ws.iter_rows() for c in r if c.value is not None)
        else:
            out[k] = text_of(p)
    return out


@pytest.mark.parametrize("card", [OOO, IP], ids=["ooo", "ip"])
def test_full_set_no_placeholders(card):
    paths = build_all(sub_deal(card))
    assert len(paths) == 6
    t = texts(paths)
    for name, txt in t.items():
        assert "{{" not in txt, f"{name}: незаполненный плейсхолдер"
    assert paths["contract"].endswith(f"САГ_Договор_{NUMBER}.docx")
    # ДКП — между продавцом и конечным покупателем, салона там нет
    assert "Иванов Иван Иванович" in t["dkp"]
    assert "Автомир" not in t["dkp"] and "Петров Пётр Петрович (ИНН" not in t["dkp"]
    # Договор с клиентом в Поручении
    assert "№ 45 от «12» сентября 2026 г." in t["contract"]
    # Конечный покупатель-физлицо в закрывающих
    assert "гражданина(ки) Российской Федерации Иванов Иван Иванович" in t["act"]


def test_ooo_preamble_and_requisites():
    t = texts(build_all(sub_deal(OOO)))
    assert ("Общество с ограниченной ответственностью «Автомир» (ИНН 7701234567, "
            "КПП 770101001, ОГРН 1157746000000), в лице Генерального директора "
            "Сидорова Сидора Сидоровича, действующего на основании Устава") in t["contract"]
    assert "7701234567 / 770101001" in t["contract"]
    assert "Сидоров С.С." in t["contract"]
    assert "ООО «Автомир», ИНН 7701234567, КПП 770101001" in t["invoice"]


def test_ip_salon():
    t = texts(build_all(sub_deal(IP)))
    assert "Индивидуальный предприниматель Петров Пётр Петрович (ИНН 770112345678, ОГРНИП 315774600000000)" in t["contract"]
    assert "ИП Петров Пётр Петрович, ИНН 770112345678, г. Москва" in t["invoice"]
    assert "Петров П.П." in t["act"]


def test_ip_end_buyer():
    t = texts(build_all(sub_deal(OOO, buyer_type="ИП", buyer_inn="180804663178")))
    assert "индивидуального предпринимателя Иванов Иван Иванович" in t["act"]
    assert "индивидуального предпринимателя Иванов Иван Иванович" in t["receipt"]
    assert "ИП Иванов Иван Иванович" in t["contract"]
    assert "ИП Иванов Иван Иванович" in t["invoice"]
    assert "Индивидуальный предприниматель Иванов Иван Иванович" in t["dkp"]


def test_invoice_numbers_numeric():
    p = build_all(sub_deal(OOO))["invoice"]
    ws = openpyxl.load_workbook(p).active
    nums = {c for c in (_as_num(v) for r in ws.iter_rows() for v in [x.value for x in r] if v is not None) if c}
    assert PRICE in nums and money.commission(PRICE, PCT, True) in nums
    assert all(isinstance(ws[c].value, (int, float)) for c in ("Z22", "Z23"))


def test_direct_account_rejected():
    data = sub_deal(OOO, account_type="direct_rf")
    with pytest.raises(SubagentSetupError):
        asyncio.run(DocumentBuilder().build_contract(data, NUMBER, "01.07.2026", PCT))


def test_missing_salon_card_rejected():
    data = sub_deal(OOO, salon="ООО «Нет такого», ИНН 7799999999")
    with pytest.raises(SubagentSetupError):
        asyncio.run(DocumentBuilder().build_invoice(data, NUMBER, "01.07.2026", PCT))


def test_direct_deal_untouched():
    data = {**DEAL, "deal_type": "Прямая"}
    data["cash_amount"] = f"{order_amount(PRICE, 82.00, NUMBER):.2f}".replace(".", ",")
    p = asyncio.run(DocumentBuilder().build_contract(data, NUMBER, "01.07.2026", PCT))
    assert p.endswith(f"АГ_Договор_{NUMBER}.docx")
    assert "Субагент" not in text_of(p)


# ── salon.py ────────────────────────────────────────────────────────────

def test_client_contract_parse():
    assert salon.parse_client_contract("№ 45 от 12.09.2026") == ("45", "12.09.2026")
    assert salon.parse_client_contract("45/А-2026 от 1.9.2026") == ("45/А-2026", "01.09.2026")
    assert salon.parse_client_contract("") == ("", "")


def test_inn_from_journal():
    assert salon.inn_from_journal("ООО «Автомир», ИНН 7701234567") == "7701234567"
    assert salon.inn_from_journal("ИП Петров, ИНН 770112345678") == "770112345678"


def test_ip_card_derived_fields():
    c = salon.normalize(IP)
    assert c["org_form"] == "ИП" and c["name_short"] == "ИП Петров Пётр Петрович"
    assert c["signer_position"] == "Индивидуальный предприниматель"
    assert salon.problems(IP) == []


def test_doverennost_basis():
    c = dict(OOO, signer_basis="доверенности № 5 от 01.09.2026")
    assert salon.full_details(c).endswith("действующего на основании доверенности № 5 от 01.09.2026")


def test_pending_choice():
    salon.save(OOO)
    salon.set_pending("chat1", "7701234567")
    assert salon.get_pending("chat1")["name_short"] == "ООО «Автомир»"
    salon.clear_pending("chat1")
    assert salon.get_pending("chat1") == {}


def test_subagent_variants():
    assert doc_builder.subagent_variants() == ["v1"]
    assert doc_builder.subagent_variant() == "v1"


def test_brand_name():
    c = dict(IP, brand="Казах Авто")
    assert salon.title(c) == "Казах Авто (ИП Петров Пётр Петрович)"
    assert salon.journal_value(c) == "Казах Авто (ИП Петров Пётр Петрович), ИНН 770112345678"
    assert salon.inn_from_journal(salon.journal_value(c)) == "770112345678"
    assert "Казах Авто" not in salon.full_details(c)       # в документы не идёт
    assert salon.problems(c) == []
