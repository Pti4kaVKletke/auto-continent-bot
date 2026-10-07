# -*- coding: utf-8 -*-
"""
Сделка «Учёт КГ» (07.10.2026): авто на кыргызских номерах и учёте, оплата до
того, как известен продавец. Этап 1 — договор (без VIN и продавца) и счёт;
этап 2 — Спецификация, ДКП, расписка, акт и отчёт по шаблонам *_KG.
"""
import asyncio
from datetime import date

import openpyxl
import pytest

import doc_builder
import gsheets_service as gs
import kg_deal
from doc_builder import DocumentBuilder, KgStageError, MissingDataError, order_amount
from tests.test_documents import DEAL, NUMBER, PRICE, PCT, text_of

STAGE2_KEYS = (
    "seller_name", "seller_birth_date", "seller_address", "seller_initials",
    "seller_id_number", "seller_id_issued_by", "seller_id_issued_date", "seller_inn",
    "car_vin", "car_body_number", "tpo_number", "tpo_date",
)


def stage1_deal(**extra):
    """Сделка на этапе 1: продавца, VIN и ТПО нет."""
    d = {k: v for k, v in DEAL.items() if k not in STAGE2_KEYS}
    d.update({
        "deal_type": kg_deal.DEAL_TYPE_KG,
        "Дата договора": "01.07.2026",
        "Дата поступления": "02.07.2026",
        "Дата расчёта": "",
        "cash_amount": f"{order_amount(PRICE, 82.00, NUMBER):.2f}".replace(".", ","),
    })
    d.update(extra)
    return d


def stage2_deal(**extra):
    d = stage1_deal()
    d.update({k: DEAL[k] for k in STAGE2_KEYS if k in DEAL and not k.startswith("tpo")})
    d.update({
        "srts": "серия KG № 1234567 от 10.07.2026",
        "gos_number": "01kg123abc",
        "poa": "№ 1234 от 12.07.2026, удостоверена нотариусом Ивановым И.И., г. Бишкек",
        "Дата ДКП": "12.07.2026",
        "Дата расчёта": "15.07.2026",
        "Фактический курс": "82,80",
        "Сумма выдана (USD)": f"{order_amount(PRICE, 82.80, NUMBER):.2f}".replace(".", ","),
    })
    d.update(extra)
    return d


def run(coro):
    return asyncio.run(coro)


# ── Тип сделки и журнал ─────────────────────────────────────────────────────

def test_is_kg_values():
    assert kg_deal.is_kg({"deal_type": "Учёт КГ"})
    assert kg_deal.is_kg({"deal_type": "учет кг"})
    assert not kg_deal.is_kg({"deal_type": "Прямая"})
    assert not kg_deal.is_kg({})


def test_journal_columns_kg_places():
    """Колонки вставлены Ильёй: H «Дата спецификации», AN–AP СРТС/Госномер/Доверенность."""
    cols = gs.COLUMNS
    assert len(cols) == 70
    assert cols.index("Дата спецификации") == 7                       # H
    assert cols.index("srts") == cols.index("tpo_date") + 1            # AN
    assert cols[cols.index("srts"):cols.index("srts") + 3] == ["srts", "gos_number", "poa"]
    assert gs.EXPECTED_HEADERS[7] == "Дата спецификации"
    i = cols.index("srts")
    assert gs.EXPECTED_HEADERS[i:i + 3] == ["СРТС", "Госномер", "Доверенность"]


# ── Даты и проверки ─────────────────────────────────────────────────────────

def test_srts_date():
    assert kg_deal.srts_date("серия KG № 1234567 от 05.10.2026") == "05.10.2026"
    assert kg_deal.srts_date("KG 1234567") == ""


def test_date_order_ok():
    assert kg_deal.date_order_problems(stage2_deal()) == []


def test_date_order_dkp_before_registration():
    d = stage2_deal(**{"Дата ДКП": "08.07.2026"})   # СРТС от 10.07
    probs = kg_deal.date_order_problems(d)
    assert probs and "дата ДКП" in probs[0]


def test_date_order_settlement_before_dkp():
    d = stage2_deal(**{"Дата расчёта": "11.07.2026"})
    assert any("дата расчёта" in p for p in kg_deal.date_order_problems(d))


def test_stage2_block_before_vehicle():
    assert "Внести данные ТС" in kg_deal.stage2_block(stage1_deal())


def test_stage2_block_needs_poa_only_for_closing():
    d = stage2_deal(poa="")
    assert kg_deal.stage2_block(d) == ""
    assert "доверенност" in kg_deal.stage2_block(d, need_poa=True)


def test_overdue_warning():
    assert kg_deal.overdue_warning(stage2_deal()) == ""
    assert "истёк" in kg_deal.overdue_warning(stage2_deal(**{"Дата ДКП": "15.08.2026"}))


def test_reminder_every_5_days():
    d = stage1_deal(Статус="Подбор", **{"Номер договора": NUMBER})
    assert kg_deal.reminder_days(d, date(2026, 7, 7)) == 5
    assert kg_deal.reminder_days(d, date(2026, 7, 8)) == 0
    assert kg_deal.reminder_days(d, date(2026, 8, 1)) == 30
    assert kg_deal.reminder_days({**d, "Статус": "завершена"}, date(2026, 7, 7)) == 0
    assert "истёк" in kg_deal.reminder_text([(d, 30)])


def test_drive_names():
    vin = DEAL["car_vin"]
    assert kg_deal.drive_name("srts", NUMBER, vin, ".pdf") == f"СРТС_{vin[-6:]}.pdf"
    assert kg_deal.drive_name("poa", NUMBER, vin, ".jpg") == f"Доверенность_{vin[-6:]}.jpg"
    assert kg_deal.drive_name("id", NUMBER, vin, ".jpg") == f"Документы_сторон_{NUMBER}.jpg"


# ── Документы этапа 1 ───────────────────────────────────────────────────────

def test_contract_stage1_without_seller_and_vin():
    path = run(DocumentBuilder().build_contract(stage1_deal(), NUMBER, "01.07.2026", PCT))
    text = text_of(path)
    assert "{{" not in text
    assert DEAL["car_vin"] not in text
    assert "Петров" not in text
    assert "Белый" in text or "белый" in text
    assert "Спецификаци" in text


def test_contract_stage1_requires_color():
    with pytest.raises(MissingDataError):
        run(DocumentBuilder().build_contract(stage1_deal(car_color=""), NUMBER, "01.07.2026", PCT))


@pytest.mark.parametrize("account_type", ["corr", "direct_rf"])
def test_invoice_without_vin(account_type):
    d = stage1_deal(account_type=account_type)
    if account_type == "direct_rf":
        d.update({"bank_name": "ПАО Сбербанк", "bank_bic": "044525225",
                  "bank_corr_acc": "30101810400000000225",
                  "account_number": "40702810000000000001"})
    path = run(DocumentBuilder().build_invoice(d, NUMBER, "01.07.2026", PCT))
    wb = openpyxl.load_workbook(path)
    cells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "VIN" not in cells
    assert "цвет белый, год выпуска 2025" in cells
    assert "{{" not in cells


def test_qr_purpose_without_vin():
    p = kg_deal.qr_purpose(NUMBER, "01.07.2026", stage1_deal(), "Toyota RAV4")
    assert "VIN" not in p and "Toyota RAV4, белый, 2025" in p


# ── Документы этапа 2 ───────────────────────────────────────────────────────

def test_stage2_docs_refused_before_vehicle():
    b = DocumentBuilder()
    for coro in (b.build_spec(stage1_deal(), NUMBER, "01.07.2026"),
                 b.build_dkp(stage1_deal(), NUMBER, "01.07.2026")):
        with pytest.raises(KgStageError):
            run(coro)


def test_spec_and_dkp():
    d = stage2_deal()
    b = DocumentBuilder()
    spec = text_of(run(b.build_spec(d, NUMBER, "01.07.2026")))
    dkp = text_of(run(b.build_dkp(d, NUMBER, "01.07.2026")))
    for t in (spec, dkp):
        assert "{{" not in t
        assert "серия KG № 1234567 от 10.07.2026" in t
        assert "01KG123ABC" in t
        assert DEAL["car_vin"] in t
    assert "ТПО" not in dkp
    assert DEAL["car_vin"][-6:] in dkp                     # номер ДКП по VIN
    assert "12» июля 2026" in spec.replace("«", "")       # дата Спецификации = дата ДКП


@pytest.mark.parametrize("kind", ["receipt", "act", "report"])
def test_closing_docs(kind):
    d = stage2_deal()
    b = DocumentBuilder()
    if kind == "receipt":
        path = run(b.build_receipt(d, NUMBER, "01.07.2026", "15.07.2026", PCT, settlement_price=PRICE))
    elif kind == "act":
        path = run(b.build_act(d, NUMBER, "01.07.2026", "15.07.2026", PCT, settlement_price=PRICE))
    else:
        path = run(b.build_report(d, NUMBER, "01.07.2026", "15.07.2026", "02.07.2026",
                                  "15.07.2026", PCT, settlement_price=PRICE))
    t = text_of(path)
    assert "{{" not in t
    assert "№ 1234 от 12.07.2026" in t
    assert "01KG123ABC" in t


def test_closing_docs_need_poa():
    with pytest.raises(KgStageError):
        run(DocumentBuilder().build_act(stage2_deal(poa=""), NUMBER, "01.07.2026",
                                        "15.07.2026", PCT, settlement_price=PRICE))


def test_closing_docs_date_order():
    d = stage2_deal(**{"Дата расчёта": "11.07.2026"})
    with pytest.raises(KgStageError):
        run(DocumentBuilder().build_receipt(d, NUMBER, "01.07.2026", "11.07.2026", PCT,
                                            settlement_price=PRICE))


def test_direct_deal_untouched():
    """Прямая сделка не видит «Учёт КГ»: VIN в счёте, ДКП со своим шаблоном."""
    b = DocumentBuilder()
    path = run(b.build_invoice(dict(DEAL), NUMBER, "01.07.2026", PCT))
    wb = openpyxl.load_workbook(path)
    cells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "VIN" in cells
    dkp = text_of(run(b.build_dkp(dict(DEAL), NUMBER, "01.07.2026")))
    assert "СРТС" not in dkp


# ── Сценарий через агента (без Drive/Sheets/Telegram) ───────────────────────

import agent as agent_mod  # noqa: E402
from tests.test_subagent_flow import FakeDrive, FakeSheets  # noqa: E402

CHAT = "chat-kg"


class KgSheets(FakeSheets):
    async def update_deal(self, num, upd):
        self.updated = (num, upd)
        if num in self.existing:
            self.existing[num] = {**self.existing[num], **upd}
        return True


@pytest.fixture
def ag(monkeypatch):
    monkeypatch.setenv("SKIP_PDF", "1")
    a = agent_mod.DocumentAgent()
    a.drive, a.sheets = FakeDrive(), KgSheets()
    a._current_chat_id = CHAT
    async def _no_scans(num): return ("sub", [])
    a.list_scan_files = _no_scans
    yield a
    kg_deal.clear_new_pending(CHAT)
    kg_deal.clear_vehicle_pending(CHAT)


def test_create_contract_kg_stage1(ag):
    kg_deal.set_new_pending(CHAT)
    data = {**DEAL, "car_price": str(PRICE), "dkp_date": "01.07.2026"}   # LLM «нашла» лишнее
    res = asyncio.run(ag._execute_tool("create_contract",
                                       {"commission_pct": 1.0, "data": data,
                                        "contract_date": "01.07.2026"}))
    assert res.get("filename", "").startswith("АГ_Договор_"), res
    names = res["extra_names"]
    assert not any(n.startswith("ДКП") for n in names)              # ДКП на этапе 1 нет
    assert any(n.startswith("Счёт") for n in names)
    saved = ag.sheets.saved["deal_data"]
    assert saved["deal_type"] == "Учёт КГ"
    for k in ("seller_name", "car_vin", "tpo_number", "dkp_date"):
        assert k not in saved
    assert not kg_deal.get_new_pending(CHAT)                          # выбор использован


def test_create_contract_kg_needs_color(ag):
    kg_deal.set_new_pending(CHAT)
    data = {**DEAL, "car_price": str(PRICE), "car_color": ""}
    res = asyncio.run(ag._execute_tool("create_contract", {"commission_pct": 1.0, "data": data}))
    assert "цвет" in res.get("error", "") and ag.sheets.saved is None


def _journal_deal(**extra):
    d = stage1_deal(**{"Номер договора": NUMBER, "Сумма Договора": "", "Платежи": "",
                       "Комиссия %": "1", "Статус": "активна"})
    d.update(extra)
    return d


def test_full_payment_sets_podbor(ag):
    total = PRICE + PRICE * PCT / 100
    ag.sheets.existing = {NUMBER: _journal_deal(**{"Дата поступления": ""})}
    res = asyncio.run(ag.add_payment_impl(NUMBER, round(total, 2), "02.07.2026"))
    assert ag.sheets.existing[NUMBER]["Статус"] == "Подбор", res
    assert any("kg_vehicle" in b["callback_data"] for b in res["buttons"])


def test_set_kg_vehicle(ag):
    ag.sheets.existing = {NUMBER: _journal_deal(Статус="Подбор")}
    data = {k: DEAL[k] for k in ("seller_name", "seller_initials", "seller_birth_date",
                                 "seller_address", "seller_id_number", "seller_id_issued_by",
                                 "seller_id_issued_date", "seller_inn", "car_vin")}
    data.update({"srts": "серия KG № 1234567 от 10.07.2026", "gos_number": "01 kg 123abc",
                 "dkp_date": "12.07.2026"})
    res = asyncio.run(ag.set_kg_vehicle_impl(NUMBER, data, chat_id=CHAT))
    deal = ag.sheets.existing[NUMBER]
    assert "внесены" in res["message"], res
    assert deal["Статус"] == "завершена"
    assert deal["Дата спецификации"] == "12.07.2026"
    assert deal["gos_number"] == "01KG123ABC"
    assert "Доверенность пока не внесена" in res["message"]


def test_set_kg_vehicle_refuses_bad_dates(ag):
    ag.sheets.existing = {NUMBER: _journal_deal(Статус="Подбор")}
    data = {k: DEAL[k] for k in ("seller_name", "seller_initials", "seller_birth_date",
                                 "seller_address", "seller_id_number", "seller_id_issued_by",
                                 "seller_id_issued_date", "car_vin")}
    data.update({"srts": "серия KG № 1234567 от 10.07.2026", "gos_number": "01KG123ABC",
                 "dkp_date": "05.07.2026"})                    # ДКП раньше регистрации
    res = asyncio.run(ag.set_kg_vehicle_impl(NUMBER, data, chat_id=CHAT))
    assert "error" in res and ag.sheets.existing[NUMBER]["Статус"] == "Подбор"


def test_scan_status_kg_counts_spec():
    names = ["Подп_АГ_Договор_1.pdf", "Подп_ДКП_ТС_1.pdf", "Подп_Расписка_1.pdf",
             "Подп_Акт_1.pdf", "Подп_Отчет_агента_1.pdf", "СРТС_066483.pdf"]
    assert agent_mod.scan_status_text(names).startswith("5/5 ✓")
    st = agent_mod.scan_status_text(names, kg=True)
    assert st.startswith("5/6") and "спецификация" in st
    assert agent_mod.scan_status_text(names + ["Подп_Спец_1.pdf"], kg=True).startswith("6/6 ✓")
