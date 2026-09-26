# -*- coding: utf-8 -*-
"""
Сборка настоящих документов по шаблонам (без Drive/Telegram): акт, отчёт,
расписка, агентский договор, счёт — для обоих вариантов бланков (v1, v2),
при полной и неполной оплате, для физлица и ИП. Проверяется, что в документ
попали правильные цифры и не осталось незаполненных {{ПЛЕЙСХОЛДЕРОВ}}.
"""
import asyncio
import re

import docx
import openpyxl
import pytest

import doc_builder
import memory
import money
from doc_builder import DocumentBuilder, order_amount

VARIANTS = doc_builder.template_variants()   # сейчас ['v1', 'v2']

PRICE = 4633490.0
PCT = 1.0
NUMBER = "010726001"   # 01.07.2026 — старое правило округления, как у реальных сделок

DEAL = {
    "buyer_name": "Иванов Иван Иванович", "buyer_birth_date": "01.01.1980",
    "buyer_address": "г. Москва, ул. Тестовая, д. 1", "buyer_initials": "Иванов И.И.",
    "passport_series": "4500", "passport_number": "123456",
    "passport_issued_by": "ОВД Тестовое", "passport_issued_date": "01.01.2005",
    "passport_code": "770-001",
    "seller_name": "Петров Пётр Петрович", "seller_birth_date": "02.02.1985",
    "seller_address": "г. Бишкек, ул. Тестовая, д. 2", "seller_initials": "Петров П.П.",
    "seller_id_number": "ID1234567", "seller_id_issued_by": "МКК КР",
    "seller_id_issued_date": "03.03.2015", "seller_inn": "12345678901234",
    "car_model": "Toyota Rav4", "car_year": "2025", "car_vin": "LBV21GT02T1C66483",
    "car_color": "Белый", "tpo_number": "41714106/220526/0000047349/00", "tpo_date": "22.05.2026",
    "car_body_number": "LBV21GT02T1C66483",
    "car_price": str(PRICE), "currency": "рублей", "cash_currency": "долларов США",
    "exchange_rate": "82,00", "Фактический курс": "82,80",
    "Дата поступления": "01.07.2026", "Дата расчёта": "05.07.2026",
    "account_type": "corr", "account_number": "1290000000123456", "account_currency": "RUB",
    "bank_name": 'ОАО "Оптима Банк"', "bank_bic": "128001", "bank_corr_acc": "30111810400000000123",
    "bank_swift": "ENEJKG22",
    "corr_bank_name": 'АО "Банк-корреспондент"', "corr_bank_bic": "044525000",
    "corr_bank_acc": "30111810900000000456",
}


def fmt(v):
    return doc_builder._fmt_num(v)


def text_of(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    for s in d.sections:
        for p in list(s.header.paragraphs) + list(s.footer.paragraphs):
            parts.append(p.text)
    return "\n".join(parts)


@pytest.fixture(params=VARIANTS)
def variant(request):
    memory.set_setting(doc_builder.TEMPLATE_VARIANT_KEY, request.param)
    yield request.param
    memory.set_setting(doc_builder.TEMPLATE_VARIANT_KEY, "v1")


def closing(received, **extra):
    """Собирает расписку, акт и отчёт так же, как бот при неполной/полной оплате."""
    legacy = money.is_legacy(NUMBER)
    total = money.add(PRICE, money.commission(PRICE, PCT, legacy), legacy)
    base = PRICE if received >= total - 0.01 else money.base_from_received(received, PCT, legacy)
    data = {**DEAL, **extra}
    data["Сумма выдана (USD)"] = f"{order_amount(base, 82.80, NUMBER):.2f}".replace(".", ",")
    data["cash_amount"] = f"{order_amount(PRICE, 82.00, NUMBER):.2f}".replace(".", ",")
    b = DocumentBuilder()

    async def run():
        return {
            "receipt": await b.build_receipt(data, NUMBER, "01.07.2026", "05.07.2026", PCT,
                                             settlement_price=base),
            "act": await b.build_act(data, NUMBER, "01.07.2026", "05.07.2026", PCT,
                                     settlement_price=base),
            "report": await b.build_report(data, NUMBER, "01.07.2026", "05.07.2026",
                                           "01.07.2026", "05.07.2026", PCT, settlement_price=base),
        }
    docs = asyncio.run(run())
    return base, data, {k: text_of(v) for k, v in docs.items()}


def test_closing_full_payment(variant):
    base, data, t = closing(round(PRICE * 1.01, 2))
    assert base == PRICE
    com = money.commission(PRICE, PCT, True)
    usd = data["Сумма выдана (USD)"].replace(".", ",")
    for name, txt in t.items():
        assert "{{" not in txt, f"{variant}/{name}: незаполненный плейсхолдер"
        assert fmt(usd.replace(",", ".")) in txt, f"{variant}/{name}: нет суммы в USD"
    assert fmt(PRICE) in t["act"] and fmt(com) in t["act"]
    assert "C66483" in t["act"]                  # номер ДКП = 6 последних знаков VIN


def test_closing_partial_payment(variant):
    base, data, t = closing(4577000.0)
    assert base == 4531683.17
    com = money.commission(base, PCT, True)
    assert com == 45316.83
    for name, txt in t.items():
        assert "{{" not in txt, f"{variant}/{name}: незаполненный плейсхолдер"
    # в акте: база и комиссия от фактически поступивших денег, а не от цены ДКП
    assert "4 531 683,17" in t["act"] and "45 316,83" in t["act"]
    assert "4 531 683,17" in t["receipt"]
    assert "54 730,47" in t["receipt"]           # 4 531 683,17 / 82,80


def test_closing_ip_buyer(variant):
    _, _, t = closing(round(PRICE * 1.01, 2), buyer_type="ИП", buyer_inn="180804663178")
    assert "Индивидуальный предприниматель Иванов Иван Иванович" in t["act"]
    assert "180804663178" in t["act"]
    assert "Гражданин(ка) Российской Федерации Иванов" not in t["act"]


def _contract(**extra):
    b = DocumentBuilder()
    data = {**DEAL, **extra}
    data["cash_amount"] = f"{order_amount(PRICE, 82.00, NUMBER):.2f}".replace(".", ",")
    path = asyncio.run(b.build_contract(data, NUMBER, "01.07.2026", PCT))
    return text_of(path)


def test_contract_individual(variant):
    t = _contract()
    assert "{{" not in t
    assert "Гражданин(ка) Российской Федерации Иванов Иван Иванович" in t
    assert "ИНН 18080" not in t


def test_contract_ip_with_bank_details(variant):
    t = _contract(buyer_type="ИП", buyer_inn="180804663178",
                  buyer_bank_details="р/с 40802810000000000001\nАО «Альфа-Банк»")
    assert "{{" not in t
    assert "ИП Иванов Иван Иванович" in t
    assert "ИНН 180804663178" in t
    assert "р/с 40802810000000000001" in t


@pytest.mark.parametrize("acc", ["corr", "direct_rf"])
def test_invoice_totals(variant, acc):
    data = {**DEAL, "account_type": acc}
    if acc == "direct_rf":
        data.update(account_number="40702810100000123456", bank_name='АО "ТБанк"',
                    bank_bic="044525974", bank_corr_acc="30101810145250000974")
    path = asyncio.run(DocumentBuilder().build_invoice(data, NUMBER, "01.07.2026", commission_pct=PCT))
    ws = openpyxl.load_workbook(path).active
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    joined = "\n".join(cells)
    assert "{{" not in joined
    # В счёте две строки (цена и комиссия), итог считает Excel формулой =SUM,
    # а сумма прописью/цифрами стоит в строке «Всего наименований 2, на сумму …».
    com = money.commission(PRICE, PCT, True)
    nums = {c for c in (_as_num(v) for v in cells) if c is not None}
    assert PRICE in nums and com in nums, "в счёте нет цены или комиссии"
    assert "4 679 824" in joined, "в счёте нет итоговой суммы 4 679 824,90"


def _as_num(v):
    try:
        return round(float(str(v).replace(" ", "").replace("\u00a0", "").replace(",", ".")), 2)
    except ValueError:
        return None


# ── Известные недочёты в документах: тесты помечены xfail и начнут «проходить»,
#    когда недочёт исправят. Тогда пометку xfail нужно снять.

@pytest.mark.xfail(reason="Марка авто нормализуется как ФИО: «BMW X3» → «Bmw X3»")
def test_known_car_model_case():
    _, _, t = closing(round(PRICE * 1.01, 2), car_model="BMW X3")
    assert "BMW X3" in t["act"]


@pytest.mark.xfail(reason="Процент комиссии в акте печатается с точкой: «комиссия 1.0%»")
def test_known_commission_pct_format():
    _, _, t = closing(round(PRICE * 1.01, 2))
    assert "1.0%" not in t["act"]


@pytest.mark.xfail(reason="Итог в счёте печатается с точкой: «на сумму 4 679 824.90 RUB»")
def test_known_invoice_total_decimal_point():
    path = asyncio.run(DocumentBuilder().build_invoice(dict(DEAL), NUMBER, "01.07.2026", commission_pct=PCT))
    ws = openpyxl.load_workbook(path).active
    joined = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "4 679 824,90" in joined
